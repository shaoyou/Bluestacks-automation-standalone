from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter


rows = [
    ("1-1", 2.06, 255, 255),
    ("1-2", 2.63, 255, 510),
    ("1-3", 3.19, 255, 765),
    ("1-4", 3.75, 255, 1020),
    ("2-1", 6.19, 258, 1278),
    ("2-2", 7.88, 258, 1536),
    ("2-3", 9.57, 258, 1794),
    ("2-4", 11.25, 258, 2052),
    ("3-1", 11.68, 265, 2317),
    ("3-2", 14.88, 265, 2582),
    ("3-3", 18.09, 265, 2847),
    ("3-4", 21.28, 265, 3112),
    ("4-1", 18.57, 273, 3385),
    ("4-2", 23.63, 273, 3658),
    ("4-3", 28.68, 273, 3931),
    ("4-4", 33.77, 273, 4204),
    ("5-1", 26.69, 280, 4484),
    ("5-2", 33.98, 280, 4764),
    ("5-3", 41.24, 280, 5044),
    ("5-4", 48.56, 280, 5324),
    ("6-1", 35.22, 288, 5612),
    ("6-2", 44.31, 288, 5900),
    ("6-3", 54.42, 288, 6188),
    ("6-4", 64.06, 288, 6476),
    ("7-1", 44.13, 295, 6771),
    ("7-2", 56.16, 295, 7066),
    ("7-3", 68.53, 295, 7361),
    ("7-4", 80.39, 295, 7656),
    ("8-1", 53.62, 303, 7959),
    ("8-2", 68.09, 303, 8262),
    ("8-3", 82.68, 303, 8565),
    ("8-4", 97.36, 303, 8868),
    ("9-1", 63.24, 310, 9178),
    ("9-2", 80.51, 310, 9488),
    ("9-3", 97.75, 310, 9798),
    ("9-4", 115.05, 310, 10108),
    ("10-1", 73.44, 318, 10426),
    ("10-2", 93.43, 318, 10744),
    ("10-3", 113.53, 318, 11062),
    ("10-4", 133.50, 318, 11380),
    ("11-1", 84.01, 325, 11705),
    ("11-2", 106.91, 325, 12030),
    ("11-3", 130.02, 325, 12355),
    ("11-4", 152.75, 325, 12680),
    ("12-1", 95.00, 333, 13013),
    ("12-2", 120.94, 333, 13346),
    ("12-3", 147.03, 333, 13679),
    ("12-4", 172.76, 333, 14012),
    ("13-1", 106.43, 340, 14352),
    ("13-2", 135.44, 340, 14692),
    ("13-3", 164.44, 340, 15032),
    ("13-4", 193.48, 340, 15372),
    ("14-1", 118.25, 348, 15720),
    ("14-2", 150.49, 348, 16068),
    ("14-3", 182.76, 348, 16416),
    ("14-4", 214.99, 348, 16764),
    ("15-1", 130.61, 355, 17119),
    ("15-2", 166.25, 355, 17474),
    ("15-3", 201.90, 355, 17829),
    ("15-4", 237.49, 355, 18184),
    ("16-1", 143.59, 363, 18547),
    ("16-2", 182.69, 363, 18910),
    ("16-3", 221.85, 363, 19273),
    ("16-4", 261.01, 363, 19636),
    ("17-1", 157.02, 370, 20006),
    ("17-2", 199.84, 370, 20376),
    ("17-3", 242.66, 370, 20746),
    ("17-4", 285.60, 370, 21116),
    ("18-1", 171.05, 378, 21494),
    ("18-2", 217.68, 378, 21872),
    ("18-3", 264.39, 378, 22250),
    ("18-4", 310.98, 378, 22628),
    ("19-1", 185.62, 385, 23013),
    ("19-2", 236.25, 385, 23398),
    ("19-3", 286.85, 385, 23783),
]

wb = Workbook()
ws = wb.active
ws.title = "关卡数据"
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A2"

headers = ["关卡", "BOSS血量(亿)", "公会币", "累计公会币"]
header_fill = PatternFill("solid", fgColor="F6C58E")
body_fill = PatternFill("solid", fgColor="FFF8F1")
first_col_fill = PatternFill("solid", fgColor="FFFFFF")
orange_font = Font(color="E46C0A", bold=True, size=16)
font_name = "PingFang SC"
header_font = Font(name=font_name, color="FFFFFF", bold=True, size=16)
body_font = Font(name=font_name, color="111111", bold=True, size=18)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")

for col, value in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col, value=value)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center
    cell.border = border

for r_idx, row in enumerate(rows, start=2):
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.alignment = center
        cell.border = border
        if c_idx == 1:
            cell.fill = first_col_fill
            cell.font = orange_font
        else:
            cell.fill = body_fill if c_idx in (2, 4) else PatternFill("solid", fgColor="FFFFFF")
            cell.font = body_font
            if c_idx == 2:
                cell.number_format = "0.00"
            else:
                cell.number_format = "0"

ws.row_dimensions[1].height = 34
for r in range(2, len(rows) + 2):
    ws.row_dimensions[r].height = 38

widths = {1: 14, 2: 18, 3: 14, 4: 18}
for idx, width in widths.items():
    ws.column_dimensions[get_column_letter(idx)].width = width

ws.auto_filter.ref = f"A1:D{len(rows) + 1}"

out = "boss_table.xlsx"
wb.save(out)
print(out)
